from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

import duckdb
import pytest
from openpyxl import Workbook, load_workbook

from wikidisputes_ssot import annotation

HEADERS = [
    "dispute_sequence",
    "dispute_id",
    "dispute_label",
    "escalated",
    "utterance_order",
    "substantive_order",
    "utterance_role",
    "utterance_id",
    "original_utterance_id",
    "speaker_id",
    "timestamp",
    "reply_to_utterance_id",
    "reply_to_utterance_id_raw",
    "reply_to_utterance_order",
    "utterance_type",
    "source_page_title",
    "utterance_text",
    "wikipedia_revision_url",
    "dispute_resolution_url",
    "utterance_text_source",
]


def _gold_row(sequence: str, role: str, order: int, uid: str, text: str) -> list[object]:
    values: dict[str, object] = {
        "dispute_sequence": sequence,
        "dispute_id": f"dispute-{sequence}",
        "dispute_label": sequence,
        "escalated": 0,
        "utterance_order": order,
        "substantive_order": None if role == "context" else order - 1,
        "utterance_role": role,
        "utterance_id": uid,
        "original_utterance_id": uid,
        "speaker_id": "speaker",
        "timestamp": "2001-01-01T00:00:00+00:00",
        "utterance_type": "original",
        "source_page_title": sequence,
        "utterance_text": text,
        "utterance_text_source": "source",
    }
    return [values.get(header) for header in HEADERS]


ROWS = {
    "d1c": _gold_row("D01", "context", 1, "d1-context", "context one"),
    "d1u2": _gold_row("D01", "utterance", 2, "d1-u2", "old two"),
    "d1u3": _gold_row("D01", "utterance", 3, "d1-u3", "old three"),
    "d2c": _gold_row("D02", "context", 1, "d2-context", "context two"),
    "d2u2": _gold_row("D02", "utterance", 2, "d2-u2", "old four"),
    "d2u3": _gold_row("D02", "utterance", 3, "d2-u3", "old five"),
}


def _write_gold(path: Path, order: list[str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Gold_Annotation"
    sheet.append(HEADERS)
    for key in order:
        sheet.append(ROWS[key])
    workbook.save(path)


def _read_gold(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet = workbook["Gold_Annotation"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    return [
        dict(zip(headers, row, strict=True)) for row in sheet.iter_rows(min_row=2, values_only=True)
    ]


def test_substantive_order_is_creation_first_and_unresolved_fallback() -> None:
    """Canonical display order keeps known chronology and unresolved bounds."""

    rows = duckdb.sql(
        f"""
        SELECT source_row_uid, ROW_NUMBER() OVER (
            ORDER BY {annotation._substantive_order_clause()}
        ) AS substantive_order
        FROM (VALUES
            ('known-early', 0, 1, 0, 'a'),
            ('unknown-between', NULL, 2, 1, 'b'),
            ('known-tie-first', 1, 3, 2, 'c'),
            ('known-tie-second', 1, 4, 2, 'd'),
            ('known-late', 2, 5, 4, 'e')
        ) AS t(
            source_row_uid,
            ssot_chronology_rank,
            canonical_display_utterance_order,
            join_display_order,
            source_order
        )
        """
    ).fetchall()

    assert [row[0] for row in rows] == [
        "known-early",
        "unknown-between",
        "known-tie-first",
        "known-tie-second",
        "known-late",
    ]
    assert [row[1] for row in rows] == [1, 2, 3, 4, 5]


def test_gold_split_part_index_precedes_display_fallback_for_tied_substantive_order() -> None:
    workbook = Workbook()
    sheet = workbook.active
    headers = [
        "dispute_sequence",
        "substantive_order",
        "utterance_role",
        "utterance_order",
        "utterance_id",
    ]
    sheet.append(headers)
    sheet.append(["D19", 9, "utterance", 9, "still-first"])
    sheet.append(["D19", 9, "utterance", 9, "belchfire"])
    sheet.append(["D19", 9, "utterance", 9, "still-last"])

    annotation._sort_gold_rows(
        sheet,
        headers,
        {2: 30, 3: 10, 4: 20},
        {2: 1, 3: 2, 4: 3},
    )

    assert [sheet.cell(row, 5).value for row in range(2, 5)] == [
        "still-first",
        "belchfire",
        "still-last",
    ]


def test_turn_integrity_fallback_uses_authoritative_nonblank_source_text(
    tmp_path: Path, monkeypatch
) -> None:
    decisions = tmp_path / "turn_integrity_decisions.parquet"
    statuses = tmp_path / "dispute_annotation_status.parquet"
    duckdb.sql(
        """
        COPY (
            SELECT
                'source-1' AS source_row_uid,
                'wikidisputes_fallback' AS final_disposition,
                '[]' AS derived_units_json,
                '' AS exclusion_reason,
                'case-1' AS case_id,
                '{}' AS detector_evidence,
                'reconstruction_rejected_wikidisputes_fallback' AS decision_reason,
                'wikidisputes_fallback' AS annotation_representation,
                'source_record_json_exact.text' AS annotation_text_source,
                'authoritative source text' AS fallback_text,
                'source_record_json_exact.text' AS fallback_text_source
        ) TO ? (FORMAT PARQUET)
        """,
        params=[str(decisions)],
    )
    duckdb.sql(
        """
        COPY (
            SELECT
                CAST(NULL AS VARCHAR) AS episode_uid,
                CAST(NULL AS VARCHAR) AS annotation_status,
                CAST(NULL AS VARCHAR) AS exclusion_reason
            WHERE FALSE
        ) TO ? (FORMAT PARQUET)
        """,
        params=[str(statuses)],
    )
    annotation_csv = tmp_path / "annotation.csv"
    with annotation_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "dispute_sequence",
                "substantive_order",
                "ssot_episode_uid",
                "ssot_source_row_uid",
                "utterance_text",
                "ssot_source_text_exact",
                "ssot_annotation_text_source",
                "ssot_text_differs_from_source",
            ],
        )
        writer.writeheader()
        writer.writerow(
            {
                "dispute_sequence": "D01",
                "substantive_order": "1",
                "ssot_episode_uid": "episode-1",
                "ssot_source_row_uid": "source-1",
                "utterance_text": "unsafe reconstructed text",
                "ssot_source_text_exact": "authoritative source text",
                "ssot_annotation_text_source": "method_b",
                "ssot_text_differs_from_source": "true",
            }
        )
        writer.writerow(
            {
                "dispute_sequence": "D01",
                "substantive_order": "2",
                "ssot_episode_uid": "episode-1",
                "ssot_source_row_uid": "source-2",
                "utterance_text": "selected text",
                "ssot_source_text_exact": "source text",
                "ssot_annotation_text_source": "method_b",
                "ssot_text_differs_from_source": "false",
            }
        )

    monkeypatch.setattr(annotation, "TURN_INTEGRITY_DECISIONS", decisions)
    monkeypatch.setattr(annotation, "DISPUTE_ANNOTATION_STATUS", statuses)
    report = annotation._turn_integrity_overlay(annotation_csv)

    assert report["included_blank_rows"] == 0
    assert report["wikidisputes_fallback_blank_rows"] == 0
    with annotation_csv.open(encoding="utf-8", newline="") as handle:
        row, kept = list(csv.DictReader(handle))
    assert row["utterance_text"] == "authoritative source text"
    assert row["ssot_annotation_text_source"] == "source_record_json_exact.text"
    assert row["ssot_text_differs_from_source"] == "false"
    assert kept["utterance_text"] == "selected text"
    assert kept["ssot_text_differs_from_source"] == "true"


def test_method_b_selection_recalculates_text_difference() -> None:
    row = {
        "utterance_text": "source text",
        "ssot_source_text_exact": "source text",
        "ssot_annotation_text_source": "source_record_json_exact.text",
        "ssot_text_differs_from_source": "false",
    }

    assert annotation._apply_method_b_text_selection(row, ("method_b", "recovered text"))
    assert row["utterance_text"] == "recovered text"
    assert row["ssot_text_differs_from_source"] == "true"

    assert annotation._apply_method_b_text_selection(row, ("method_b", "source text"))
    assert row["utterance_text"] == "source text"
    assert row["ssot_text_differs_from_source"] == "false"


def test_audited_gold_units_keep_membership_and_explicit_review() -> None:
    root = Path(__file__).resolve().parents[1]
    annotation_csv = root / "output/annotation/wikidisputes_llm_annotation_input.csv"
    staged_csv = root / "output/annotation/wikidisputes_llm_annotation_input.method_b_staged.csv"
    gold_path = root / "output/annotation/gold_input_ssot_annotation_ready.xlsx"
    if not all(path.exists() for path in (annotation_csv, staged_csv, gold_path)):
        pytest.skip("rebuilt annotation and Gold artifacts are required")

    retained = Counter(
        [("D00003", 2), ("D00003", 6), ("D00111", 4), ("D00181", 22), ("D00977", 5)]
        + [("D01057", order) for order in (3, 6, 9, 13, 18, 21, 22, 24, 26, 28, 32)]
        + [("D01057", 11)] * 3
        + [("D01342", 17), ("D03503", 9), ("D03503", 17)]
        + [("D03977", order) for order in (4, 5, 12, 13)]
        + [("D05465", 8), ("D05549", 13), ("D05549", 27), ("D06530", 11)]
    )
    assert sum(retained.values()) == 30
    workbook = load_workbook(gold_path, read_only=True, data_only=True)
    try:
        records = workbook["Gold_Annotation"].values
        headers = next(records)
        gold_rows = [dict(zip(headers, values, strict=True)) for values in records]
    finally:
        workbook.close()
    gold_membership = Counter(
        (str(row["dispute_sequence"]), int(row["substantive_order"])) for row in gold_rows
    )
    assert all(gold_membership[key] >= count for key, count in retained.items())

    audited_ids = {
        "181458312.3149.3149",
        "260624906.75956.75956",
        "308287895.17875.17875",
        "715133338.29482.29482",
        "133430871.30166.30166",
        "133975276.36783.36783",
        "606024046.22544.22544",
        "662046446.113863.113863",
    }

    def selected_rows(path: Path) -> dict[str, dict[str, str]]:
        with path.open(encoding="utf-8", newline="") as handle:
            return {
                row["utterance_id"]: row
                for row in csv.DictReader(handle)
                if row["utterance_id"] in audited_ids
            }

    exported = selected_rows(annotation_csv)
    staged = selected_rows(staged_csv)
    assert "308287895.17875.17875" not in exported
    assert set(exported) == audited_ids - {"308287895.17875.17875"}
    assert exported["181458312.3149.3149"]["speaker_id"] == "BlastOButter42"
    assert (
        exported["181458312.3149.3149"]["utterance_text"]
        == staged["181458312.3149.3149"]["utterance_text"]
    )
    for utterance_id in audited_ids - {"308287895.17875.17875", "181458312.3149.3149"}:
        assert exported[utterance_id]["utterance_text"] == staged[utterance_id]["utterance_text"]
        assert exported[utterance_id]["speaker_id"] == staged[utterance_id]["speaker_id"]
        assert exported[utterance_id]["needs_rereview"] == "true"
    gold_by_id = {str(row["utterance_id"]): row for row in gold_rows}
    assert "308287895.17875.17875" not in gold_by_id
    for utterance_id in audited_ids - {"308287895.17875.17875"}:
        assert gold_by_id[utterance_id]["provenance"] == "needs_rereview"


def test_full_export_keeps_chronology_and_display_fields_separate() -> None:
    query = annotation.full_export_sql()

    assert "o.ssot_chronology_rank AS utterance_order" in query
    assert "o.canonical_display_utterance_order AS ssot_display_utterance_order" in query
    assert "o.local_display_order AS display_order" in query
    assert "o.local_substantive_order AS substantive_order" in query
    assert "join_display_order NULLS LAST" in query


def test_gold_export_canonicalizes_physical_order_deterministically(
    tmp_path: Path, monkeypatch
) -> None:
    annotation_csv = tmp_path / "annotation.csv"
    fieldnames = [
        "dispute_id",
        "utterance_id",
        "utterance_role",
        "ssot_source_row_uid",
        "ssot_context_node_uid",
        "utterance_text",
        "utterance_order",
        "display_order",
    ]
    with annotation_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        rows = (
            ("D01", "d1-context", "context", "context one", "ctx-d1", "", 1),
            ("D01", "d1-u2", "utterance", "", "", 1, 2),
            ("D01", "d1-u3", "utterance", "", "", 2, 3),
            ("D02", "d2-context", "context", "context two", "ctx-d2", "", 1),
            ("D02", "d2-u2", "utterance", "", "", 1, 2),
            ("D02", "d2-u3", "utterance", "", "", 2, 3),
        )
        for sequence, uid, role, text, context_uid, rank, display_order in rows:
            writer.writerow(
                {
                    "dispute_id": f"dispute-{sequence}",
                    "utterance_id": uid,
                    "utterance_role": role,
                    "ssot_source_row_uid": f"source-{uid}",
                    "ssot_context_node_uid": context_uid,
                    "utterance_text": text,
                    "utterance_order": rank,
                    "display_order": display_order,
                }
            )

    selection = tmp_path / "selection.parquet"
    duckdb.sql(
        """
        COPY (
            SELECT * FROM (VALUES
                ('source-d1-u2', 'method_a', 'selected two'),
                ('source-d1-u3', 'method_b', 'selected three'),
                ('source-d2-u2', 'method_a_fallback', 'selected four'),
                ('source-d2-u3', 'method_a', 'selected five')
            ) AS t(source_row_uid, selected_method, selected_text)
        ) TO ? (FORMAT PARQUET)
        """,
        params=[str(selection)],
    )

    monkeypatch.setattr(annotation, "ANNOTATION", tmp_path / "output")
    monkeypatch.setattr(annotation, "FINAL_SELECTION", selection)
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    _write_gold(first, ["d2u3", "d1u3", "d2c", "d1c", "d2u2", "d1u2"])
    _write_gold(second, ["d1u2", "d2c", "d2u2", "d1u3", "d1c", "d2u3"])
    enriched = load_workbook(first)
    enriched_sheet = enriched["Gold_Annotation"]
    enriched.create_sheet("Codebook")
    for column in range(21, 48):
        enriched_sheet.cell(1, column, f"coding_column_{column}")
        enriched_sheet.cell(2, column, "ignored")
    enriched.save(first)

    first_report = annotation.export_annotation_ready_gold(first, annotation_csv)
    output = annotation.ANNOTATION / annotation.FINAL_GOLD_NAME
    first_hash = hashlib.sha256(output.read_bytes()).hexdigest()
    first_rows = _read_gold(output)
    second_report = annotation.export_annotation_ready_gold(second, annotation_csv)
    second_hash = hashlib.sha256(output.read_bytes()).hexdigest()
    second_rows = _read_gold(output)

    final_workbook = load_workbook(output, read_only=True, data_only=False)
    assert final_workbook.sheetnames == ["Gold_Annotation"]
    assert final_workbook["Gold_Annotation"].max_column == 21
    final_workbook.close()

    assert first_hash == second_hash
    assert first_rows == second_rows
    assert first_report["context_rows"] == second_report["context_rows"] == 0
    assert first_report["context_classified_rows"] == 2
    assert second_report["context_classified_rows"] == 2
    assert [row["utterance_id"] for row in first_rows] == [
        "d1-context",
        "d1-u2",
        "d1-u3",
        "d2-context",
        "d2-u2",
        "d2-u3",
    ]

    for sequence in ("D01", "D02"):
        dispute = [row for row in first_rows if row["dispute_sequence"] == sequence]
        orders = [
            int(row["utterance_order"]) for row in dispute if row["utterance_order"] is not None
        ]
        assert orders == sorted(orders)

    assert Counter(
        (row["utterance_id"], row["utterance_text"], row["provenance"]) for row in first_rows
    ) == Counter(
        {
            ("d1-context", "context one", "wikidisputes_source"): 1,
            ("d1-u2", "selected two", "method_a"): 1,
            ("d1-u3", "selected three", "method_b"): 1,
            ("d2-context", "context two", "wikidisputes_source"): 1,
            ("d2-u2", "selected four", "method_a_fallback"): 1,
            ("d2-u3", "selected five", "method_a"): 1,
        }
    )
    assert all(row["utterance_role"] == "utterance" for row in first_rows)
    assert all("context" not in str(row["provenance"]) for row in first_rows)
    former_contexts = [row for row in first_rows if row["provenance"] == "wikidisputes_source"]
    assert all(row["utterance_order"] is None for row in former_contexts)


def test_gold_export_projects_current_dispute_identity_from_final_ssot_row(
    tmp_path: Path, monkeypatch
) -> None:
    annotation_csv = tmp_path / "annotation.csv"
    fields = [
        "dispute_sequence",
        "dispute_id",
        "dispute_label",
        "utterance_id",
        "original_utterance_id",
        "utterance_role",
        "ssot_source_row_uid",
        "utterance_text",
        "utterance_order",
        "display_order",
    ]
    with annotation_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerow(
            {
                "dispute_sequence": "D00042",
                "dispute_id": "current-id",
                "dispute_label": "Current label",
                "utterance_id": "d1-u2",
                "original_utterance_id": "d1-u2",
                "utterance_role": "utterance",
                "ssot_source_row_uid": "source-d1-u2",
                "utterance_text": "current text",
                "utterance_order": "1",
                "display_order": "1",
            }
        )

    selection = tmp_path / "selection.parquet"
    duckdb.sql(
        """
        COPY (
            SELECT * FROM (VALUES
                ('source-d1-u2', 'method_a', 'selected text')
            ) AS t(source_row_uid, selected_method, selected_text)
        ) TO ? (FORMAT PARQUET)
        """,
        params=[str(selection)],
    )
    monkeypatch.setattr(annotation, "ANNOTATION", tmp_path / "output")
    monkeypatch.setattr(annotation, "FINAL_SELECTION", selection)

    gold = tmp_path / "gold.xlsx"
    _write_gold(gold, ["d1u2"])
    workbook = load_workbook(gold)
    sheet = workbook["Gold_Annotation"]
    sheet.cell(2, HEADERS.index("dispute_sequence") + 1).value = "legacy-sequence"
    sheet.cell(2, HEADERS.index("dispute_id") + 1).value = "legacy-id"
    sheet.cell(2, HEADERS.index("dispute_label") + 1).value = "Legacy label"
    workbook.save(gold)

    report = annotation.export_annotation_ready_gold(gold, annotation_csv)
    [row] = _read_gold(Path(report["path"]))
    assert row["dispute_sequence"] == "D00042"
    assert row["dispute_id"] == "current-id"
    assert row["dispute_label"] == "Current label"


def test_gold_readiness_marks_only_changed_units_and_preserves_existing_annotations(
    tmp_path: Path, monkeypatch
) -> None:
    annotation_csv = tmp_path / "annotation.csv"
    fields = [
        "dispute_sequence",
        "dispute_id",
        "dispute_label",
        "utterance_id",
        "utterance_role",
        "ssot_source_row_uid",
        "utterance_text",
        "utterance_order",
        "display_order",
        "needs_rereview",
        "ssot_turn_integrity_case_id",
        "ssot_turn_integrity_decision_reason",
    ]
    with annotation_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for order, uid, changed in ((1, "d1-u2", False), (2, "d1-u3", True)):
            writer.writerow(
                {
                    "dispute_sequence": "D01",
                    "dispute_id": "dispute-D01",
                    "dispute_label": "D01",
                    "utterance_id": uid,
                    "utterance_role": "utterance",
                    "ssot_source_row_uid": f"source-{uid}",
                    "utterance_text": f"current {order}",
                    "utterance_order": order,
                    "display_order": order,
                    "needs_rereview": str(changed).lower(),
                    "ssot_turn_integrity_case_id": "case-changed" if changed else "",
                    "ssot_turn_integrity_decision_reason": (
                        "fragment_reattached" if changed else ""
                    ),
                }
            )
    selection = tmp_path / "selection.parquet"
    duckdb.sql(
        "COPY (SELECT * FROM (VALUES "
        "('source-d1-u2', 'method_a', 'current 1'), "
        "('source-d1-u3', 'method_a', 'current 2')"
        ") AS t(source_row_uid, selected_method, selected_text)) "
        "TO ? (FORMAT PARQUET)",
        params=[str(selection)],
    )
    monkeypatch.setattr(annotation, "ANNOTATION", tmp_path / "output")
    monkeypatch.setattr(annotation, "FINAL_SELECTION", selection)
    gold = tmp_path / "gold.xlsx"
    _write_gold(gold, ["d1u2", "d1u3"])
    workbook = load_workbook(gold)
    sheet = workbook["Gold_Annotation"]
    sheet.cell(2, HEADERS.index("escalated") + 1).value = 1
    sheet.cell(3, HEADERS.index("escalated") + 1).value = 0
    workbook.save(gold)

    report = annotation.export_annotation_ready_gold(gold, annotation_csv)
    rows = _read_gold(Path(report["path"]))
    assert [(row["utterance_id"], row["escalated"], row["provenance"]) for row in rows] == [
        ("d1-u2", 1, "method_a"),
        ("d1-u3", 0, "needs_rereview"),
    ]
    readiness = report["annotation_readiness"]
    assert readiness["annotation_ready"] is False
    assert readiness["blocking_case_count"] == 1
    assert readiness["blocking_cases"] == [
        {
            "utterance_id": "d1-u3",
            "source_row_uid": "source-d1-u3",
            "case_id": "case-changed",
            "type": "fragment_reattached",
        }
    ]

    decisions = tmp_path / "decisions.parquet"
    duckdb.sql(
        "COPY (SELECT * FROM (VALUES "
        "('source-d1-u2', 'case-identity', TRUE, 'unresolved_physical_identity')"
        ") AS t(source_row_uid, case_id, annotation_blocking, "
        "annotation_blocking_reason)) TO ? (FORMAT PARQUET)",
        params=[str(decisions)],
    )
    monkeypatch.setattr(annotation, "TURN_INTEGRITY_DECISIONS", decisions)
    monkeypatch.setattr(
        annotation, "DISPUTE_ANNOTATION_STATUS", tmp_path / "missing-status.parquet"
    )
    report = annotation.export_annotation_ready_gold(gold, annotation_csv)
    rows = _read_gold(Path(report["path"]))
    assert rows[0]["escalated"] == 1
    assert rows[0]["provenance"] == "needs_rereview"
    assert report["annotation_readiness"]["blocking_case_count"] == 2
    assert report["annotation_readiness"]["blocking_cases"][0] == {
        "utterance_id": "d1-u2",
        "source_row_uid": "source-d1-u2",
        "case_id": "case-identity",
        "type": "unresolved_physical_identity",
    }


def test_gold_split_children_inherit_current_d01057_membership_and_part_order(
    tmp_path: Path, monkeypatch
) -> None:
    annotation_csv = tmp_path / "annotation.csv"
    fields = [
        "dispute_sequence",
        "dispute_id",
        "dispute_label",
        "utterance_id",
        "original_utterance_id",
        "utterance_role",
        "ssot_source_row_uid",
        "utterance_text",
        "utterance_order",
        "substantive_order",
        "display_order",
        "ssot_turn_integrity_disposition",
        "ssot_turn_integrity_part_index",
    ]
    children = [
        ("turn-unit:v1:still-first", "Still", 1),
        ("turn-unit:v1:belchfire", "Belchfire", 2),
        ("turn-unit:v1:still-last", "Still", 3),
    ]
    with annotation_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for uid, text, part_index in children:
            writer.writerow(
                {
                    "dispute_sequence": "D01057",
                    "dispute_id": "current-d01057",
                    "dispute_label": "Current D01057 label",
                    "utterance_id": uid,
                    "original_utterance_id": "legacy-composite",
                    "utterance_role": "utterance",
                    "ssot_source_row_uid": "source-composite",
                    "utterance_text": text,
                    "utterance_order": "9",
                    "substantive_order": "9",
                    "display_order": "9",
                    "ssot_turn_integrity_disposition": "split",
                    "ssot_turn_integrity_part_index": str(part_index),
                }
            )

    selection = tmp_path / "selection.parquet"
    duckdb.sql(
        """
        COPY (
            SELECT 'source-composite' AS source_row_uid,
                   'method_a' AS selected_method,
                   'unused' AS selected_text
        ) TO ? (FORMAT PARQUET)
        """,
        params=[str(selection)],
    )
    monkeypatch.setattr(annotation, "ANNOTATION", tmp_path / "output")
    monkeypatch.setattr(annotation, "FINAL_SELECTION", selection)
    monkeypatch.setattr(annotation, "TURN_INTEGRITY_DECISIONS", tmp_path / "missing.parquet")
    monkeypatch.setattr(
        annotation, "DISPUTE_ANNOTATION_STATUS", tmp_path / "missing-status.parquet"
    )

    gold = tmp_path / "gold.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Gold_Annotation"
    sheet.append(HEADERS)
    for uid, text, _ in reversed(children):
        values = dict(zip(HEADERS, _gold_row("D19", "utterance", 9, uid, text), strict=True))
        values["dispute_id"] = "legacy-d19"
        values["dispute_label"] = "Legacy D19 label"
        sheet.append([values[header] for header in HEADERS])
    workbook.save(gold)

    report = annotation.export_annotation_ready_gold(gold, annotation_csv)
    rows = _read_gold(Path(report["path"]))

    assert report["complete_current_ssot_disputes"] == 1
    assert [row["utterance_text"] for row in rows] == ["Still", "Belchfire", "Still"]
    assert {row["dispute_sequence"] for row in rows} == {"D01057"}
    assert {row["dispute_id"] for row in rows} == {"current-d01057"}
    assert {row["dispute_label"] for row in rows} == {"Current D01057 label"}


def test_gold_export_rejects_partial_current_ssot_dispute(tmp_path: Path, monkeypatch) -> None:
    annotation_csv = tmp_path / "annotation.csv"
    with annotation_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "dispute_sequence",
                "dispute_id",
                "dispute_label",
                "utterance_id",
                "utterance_role",
                "ssot_source_row_uid",
            ],
        )
        writer.writeheader()
        for uid in ("d1-u2", "d1-u3"):
            writer.writerow(
                {
                    "dispute_sequence": "D01",
                    "dispute_id": "dispute-D01",
                    "dispute_label": "D01",
                    "utterance_id": uid,
                    "utterance_role": "utterance",
                    "ssot_source_row_uid": f"source-{uid}",
                }
            )
    selection = tmp_path / "selection.parquet"
    duckdb.sql(
        """
        COPY (
            SELECT 'source-d1-u2' AS source_row_uid,
                   'method_a' AS selected_method,
                   'selected' AS selected_text
        ) TO ? (FORMAT PARQUET)
        """,
        params=[str(selection)],
    )
    monkeypatch.setattr(annotation, "ANNOTATION", tmp_path / "output")
    monkeypatch.setattr(annotation, "FINAL_SELECTION", selection)
    monkeypatch.setattr(annotation, "TURN_INTEGRITY_DECISIONS", tmp_path / "missing.parquet")
    monkeypatch.setattr(
        annotation, "DISPUTE_ANNOTATION_STATUS", tmp_path / "missing-status.parquet"
    )
    gold = tmp_path / "gold.xlsx"
    _write_gold(gold, ["d1u2"])

    with pytest.raises(RuntimeError, match="not one complete current SSOT dispute"):
        annotation.export_annotation_ready_gold(gold, annotation_csv)


def test_gold_export_does_not_require_a_context_row(tmp_path: Path, monkeypatch) -> None:
    annotation_csv = tmp_path / "annotation.csv"
    with annotation_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "dispute_id",
                "utterance_id",
                "utterance_role",
                "ssot_source_row_uid",
            ],
        )
        writer.writeheader()
        for uid in ("d1-u2", "d1-u3"):
            writer.writerow(
                {
                    "dispute_id": "dispute-D01",
                    "utterance_id": uid,
                    "utterance_role": "utterance",
                    "ssot_source_row_uid": f"source-{uid}",
                }
            )

    selection = tmp_path / "selection.parquet"
    duckdb.sql(
        """
        COPY (
            SELECT * FROM (VALUES
                ('source-d1-u2', 'method_a', 'selected two'),
                ('source-d1-u3', 'method_b', 'selected three')
            ) AS t(source_row_uid, selected_method, selected_text)
        ) TO ? (FORMAT PARQUET)
        """,
        params=[str(selection)],
    )
    monkeypatch.setattr(annotation, "ANNOTATION", tmp_path / "output")
    monkeypatch.setattr(annotation, "FINAL_SELECTION", selection)

    gold = tmp_path / "gold.xlsx"
    _write_gold(gold, ["d1u3", "d1u2"])
    report = annotation.export_annotation_ready_gold(gold, annotation_csv)
    rows = _read_gold(Path(report["path"]))

    assert report["context_rows"] == 0
    assert [row["utterance_id"] for row in rows] == ["d1-u2", "d1-u3"]
    assert all(row["provenance"] in {"method_a", "method_b"} for row in rows)


def test_gold_export_does_not_force_context_to_first_row(tmp_path: Path, monkeypatch) -> None:
    annotation_csv = tmp_path / "annotation.csv"
    fieldnames = [
        "dispute_id",
        "utterance_id",
        "utterance_role",
        "ssot_source_row_uid",
        "display_order",
    ]
    with annotation_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for uid, role, display_order in (
            ("d1-u2", "utterance", 1),
            ("d1-context", "context", 2),
            ("d1-u3", "utterance", 3),
        ):
            writer.writerow(
                {
                    "dispute_id": "dispute-D01",
                    "utterance_id": uid,
                    "utterance_role": role,
                    "ssot_source_row_uid": f"source-{uid}",
                    "display_order": display_order,
                }
            )

    selection = tmp_path / "selection.parquet"
    duckdb.sql(
        """
        COPY (
            SELECT * FROM (VALUES
                ('source-d1-u2', 'method_a', 'selected two'),
                ('source-d1-u3', 'method_b', 'selected three')
            ) AS t(source_row_uid, selected_method, selected_text)
        ) TO ? (FORMAT PARQUET)
        """,
        params=[str(selection)],
    )
    monkeypatch.setattr(annotation, "ANNOTATION", tmp_path / "output")
    monkeypatch.setattr(annotation, "FINAL_SELECTION", selection)

    gold = tmp_path / "gold.xlsx"
    _write_gold(gold, ["d1c", "d1u3", "d1u2"])
    workbook = load_workbook(gold)
    sheet = workbook["Gold_Annotation"]
    order_by_id = {"d1-u2": 1, "d1-context": 2, "d1-u3": 3}
    for row in range(2, sheet.max_row + 1):
        utterance_id = sheet.cell(row, HEADERS.index("utterance_id") + 1).value
        sheet.cell(row, HEADERS.index("utterance_order") + 1, order_by_id[utterance_id])
    workbook.save(gold)
    report = annotation.export_annotation_ready_gold(gold, annotation_csv)
    rows = _read_gold(Path(report["path"]))

    assert [row["utterance_id"] for row in rows] == ["d1-u2", "d1-context", "d1-u3"]
    context = rows[1]
    assert context["utterance_role"] == "utterance"
    assert context["provenance"] == "wikidisputes_source"
    assert all(row["utterance_role"] == "utterance" for row in rows)
    assert all("context" not in str(row["provenance"]) for row in rows)


def test_gold_export_applies_only_configured_discussion_exclusions(
    tmp_path: Path, monkeypatch
) -> None:
    annotation_csv = tmp_path / "annotation.csv"
    with annotation_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "dispute_id",
                "dispute_label",
                "utterance_id",
                "original_utterance_id",
                "utterance_role",
                "ssot_source_row_uid",
                "ssot_episode_uid",
                "ssot_logical_utterance_uid",
            ],
        )
        writer.writeheader()
        for sequence, uid in (("D01", "d1-u2"), ("D02", "d2-u2")):
            writer.writerow(
                {
                    "dispute_id": f"dispute-{sequence}",
                    "dispute_label": sequence,
                    "utterance_id": uid,
                    "original_utterance_id": uid,
                    "utterance_role": "utterance",
                    "ssot_source_row_uid": f"source-{uid}",
                    "ssot_episode_uid": f"episode-{sequence}",
                    "ssot_logical_utterance_uid": f"logical-{uid}",
                }
            )

    selection = tmp_path / "selection.parquet"
    duckdb.sql(
        """
        COPY (
            SELECT * FROM (VALUES
                ('source-d1-u2', 'method_a', 'selected two'),
                ('source-d2-u2', 'method_a', 'selected four')
            ) AS t(source_row_uid, selected_method, selected_text)
        ) TO ? (FORMAT PARQUET)
        """,
        params=[str(selection)],
    )
    exclusions = tmp_path / "exclusions.json"
    exclusions.write_text(
        '{"exclusions":[{"dispute_id":"dispute-D02",'
        '"dispute_label":"D02","reason":"confirmed_malformed"}]}',
        encoding="utf-8",
    )
    monkeypatch.setattr(annotation, "ANNOTATION", tmp_path / "output")
    monkeypatch.setattr(annotation, "FINAL_SELECTION", selection)
    monkeypatch.setattr(annotation, "ANNOTATION_EXCLUSIONS", exclusions)

    gold = tmp_path / "gold.xlsx"
    _write_gold(gold, ["d1c", "d1u2", "d2c", "d2u2"])
    report = annotation.export_annotation_ready_gold(gold, annotation_csv)
    rows = _read_gold(Path(report["path"]))

    assert {row["dispute_id"] for row in rows} == {"dispute-D01"}
    assert report["rows"] == 1
    assert report["excluded_discussions"][0]["reason"] == "confirmed_malformed"


def test_turn_integrity_reply_resolution_preserves_raw_and_rejects_ambiguous_split() -> None:
    rows = [
        {
            "ssot_episode_uid": "episode-1",
            "dispute_sequence": "D01",
            "substantive_order": "1",
            "utterance_id": "root",
            "original_utterance_id": "root",
            "reply_to_utterance_id": "",
            "reply_to_utterance_id_raw": "",
        },
        {
            "ssot_episode_uid": "episode-1",
            "dispute_sequence": "D01",
            "substantive_order": "2",
            "utterance_id": "child-a",
            "original_utterance_id": "replayed-source",
            "reply_to_utterance_id": "",
            "reply_to_utterance_id_raw": "",
        },
        {
            "ssot_episode_uid": "episode-1",
            "dispute_sequence": "D01",
            "substantive_order": "3",
            "utterance_id": "child-b",
            "original_utterance_id": "replayed-source",
            "reply_to_utterance_id": "",
            "reply_to_utterance_id_raw": "",
        },
        {
            "ssot_episode_uid": "episode-1",
            "dispute_sequence": "D01",
            "substantive_order": "4",
            "utterance_id": "reply-to-root",
            "original_utterance_id": "reply-to-root",
            "reply_to_utterance_id": "root",
            "reply_to_utterance_id_raw": "root",
        },
        {
            "ssot_episode_uid": "episode-1",
            "dispute_sequence": "D01",
            "substantive_order": "5",
            "utterance_id": "ambiguous-reply",
            "original_utterance_id": "ambiguous-reply",
            "reply_to_utterance_id": "replayed-source",
            "reply_to_utterance_id_raw": "replayed-source",
        },
    ]

    report = annotation._resolve_overlay_replies(rows)

    assert report["resolved"] == 1
    assert report["ambiguous_split_targets"] == 1
    assert rows[4]["reply_to_utterance_id_raw"] == "replayed-source"
    assert rows[4]["reply_to_utterance_id"] == ""
    assert rows[4]["ssot_reply_resolution_status"] == "unresolved_ambiguous_split_target"
    assert rows[3]["reply_to_utterance_id"] == "root"
    assert rows[3]["reply_to_utterance_id_raw"] == "root"


def test_turn_integrity_reply_resolution_follows_transitive_suppressed_anchor() -> None:
    rows = [
        {
            "ssot_episode_uid": "episode-1",
            "dispute_sequence": "D01",
            "ssot_source_row_uid": "anchor-source",
            "substantive_order": "1",
            "utterance_id": "anchor-id",
            "original_utterance_id": "anchor-id",
            "reply_to_utterance_id": "",
            "reply_to_utterance_id_raw": "",
        },
        {
            "ssot_episode_uid": "episode-1",
            "dispute_sequence": "D01",
            "ssot_source_row_uid": "reply-source",
            "substantive_order": "2",
            "utterance_id": "reply-id",
            "original_utterance_id": "reply-id",
            "reply_to_utterance_id": "suppressed-id",
            "reply_to_utterance_id_raw": "suppressed-id",
        },
    ]

    report = annotation._resolve_overlay_replies(
        rows,
        source_anchor_map={
            "suppressed-source": {"intermediate-source"},
            "intermediate-source": {"anchor-source"},
        },
        source_alias_map={"suppressed-id": {"suppressed-source"}},
    )

    assert report["resolved"] == 1
    assert rows[1]["reply_to_utterance_id_raw"] == "suppressed-id"
    assert rows[1]["reply_to_utterance_id"] == "anchor-id"
    assert rows[1]["ssot_reply_resolution_status"] == "resolved_after_turn_integrity"


def test_unresolved_anchor_evidence_cannot_redirect_replies() -> None:
    anchors = annotation._source_anchor_map(
        {
            "unresolved-source": [
                {
                    "final_disposition": "keep",
                    "evidence": '{"anchor_source_row_uid":"retained-source"}',
                }
            ],
            "proven-source": [
                {
                    "final_disposition": "alias_or_suppress_duplicate",
                    "evidence": (
                        '{"lifecycle_identity":"proven_alias",'
                        '"anchor_source_row_uid":"retained-source"}'
                    ),
                }
            ],
        }
    )

    assert "unresolved-source" not in anchors
    assert anchors["proven-source"] == {"retained-source"}


def test_annotation_readiness_prefers_decisions_without_double_counting_summary(
    tmp_path: Path, monkeypatch
) -> None:
    decisions = tmp_path / "decisions.parquet"
    duckdb.sql(
        """
        COPY (
            SELECT * FROM (VALUES
                ('source-1', 'case-1', TRUE, 'unresolved_replay_identity', 'keep'),
                ('source-2', 'case-2', FALSE, '', 'keep')
            ) AS t(
                source_row_uid,
                case_id,
                annotation_blocking,
                annotation_blocking_reason,
                final_disposition
            )
        ) TO ? (FORMAT PARQUET)
        """,
        params=[str(decisions)],
    )
    reports = tmp_path / "reports"
    summary_path = reports / "turn_integrity" / "repair_summary.json"
    summary_path.parent.mkdir(parents=True)
    summary_path.write_text(
        json.dumps(
            {
                "annotation_blockers": {
                    "count": 99,
                    "by_reason": {"stale_summary_should_not_be_added": 99},
                }
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(annotation, "TURN_INTEGRITY_DECISIONS", decisions)
    monkeypatch.setattr(annotation, "REPORTS", reports)
    monkeypatch.setattr(
        annotation, "DISPUTE_ANNOTATION_STATUS", tmp_path / "missing-status.parquet"
    )

    readiness = annotation._annotation_readiness()

    assert readiness["status"] == "non_ready"
    assert readiness["annotation_ready"] is False
    assert readiness["blocking_case_count"] == 1
    assert readiness["blocking_by_type"] == {"unresolved_replay_identity": 1}
    assert readiness["source"] == "turn_integrity_decisions"


@pytest.mark.parametrize(
    ("payload", "expected_count", "expected_types", "expected_status"),
    [
        ({"count": 0, "by_reason": {}}, 0, {}, "pass"),
        (
            {"count": 2, "by_reason": {"unresolved_composite": 2}},
            2,
            {"unresolved_composite": 2},
            "non_ready",
        ),
        (
            [
                {
                    "source_row_uid": "source-1",
                    "annotation_blocking_reason": "speaker_signature_conflict",
                }
            ],
            1,
            {"speaker_signature_conflict": 1},
            "non_ready",
        ),
    ],
)
def test_annotation_readiness_accepts_summary_count_and_case_shapes(
    tmp_path: Path,
    monkeypatch,
    payload: object,
    expected_count: int,
    expected_types: dict[str, int],
    expected_status: str,
) -> None:
    reports = tmp_path / "reports"
    summary_path = reports / "turn_integrity" / "repair_summary.json"
    summary_path.parent.mkdir(parents=True)
    summary_path.write_text(json.dumps({"annotation_blockers": payload}), encoding="utf-8")
    monkeypatch.setattr(annotation, "TURN_INTEGRITY_DECISIONS", tmp_path / "missing.parquet")
    monkeypatch.setattr(annotation, "REPORTS", reports)
    monkeypatch.setattr(
        annotation, "DISPUTE_ANNOTATION_STATUS", tmp_path / "missing-status.parquet"
    )

    readiness = annotation._annotation_readiness()

    assert readiness["status"] == expected_status
    assert readiness["blocking_case_count"] == expected_count
    assert readiness["blocking_by_type"] == expected_types
    assert readiness["annotation_ready"] is (expected_count == 0)
